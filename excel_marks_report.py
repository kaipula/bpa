import os
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

os.system("")  # enable ANSI colours in Windows terminal
RED, RESET = "\033[91m", "\033[0m"

OUT = "marks_report.xlsx"
SUBJECTS = ["HCI", "BPA", "DP", "CC"]
PASS_MARK = 50

students = [
    ("21CS001", "Swetha", [92, 88, 95, 90]),
    ("21CS002", "Arun",   [76, 45, 68, 70]),
    ("21CS003", "Priya",  [85, 90, 82, 88]),
    ("21CS004", "Karthik", [55, 60, 48, 52]),
    ("21CS005", "Divya",  [98, 94, 91, 96]),
    ("21CS006", "Rahul",  [65, 72, 70, 61]),
]


def get_grade(avg, passed):
    if not passed:
        return "RA"
    if avg >= 90:
        return "O"
    if avg >= 80:
        return "A+"
    if avg >= 70:
        return "A"
    if avg >= 60:
        return "B+"
    return "B"


wb = Workbook()
ws = wb.active
ws.title = "Marks"
ws.append(["RegNo", "Name"] + SUBJECTS + ["Total", "Average", "Grade", "Result"])

for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", start_color="4F81BD")

passed_count = failed_count = 0
for reg, name, marks in students:
    total = sum(marks)
    avg = total / len(marks)
    passed = all(m >= PASS_MARK for m in marks)
    grade = get_grade(avg, passed)
    result = "PASS" if passed else "FAIL"
    ws.append([reg, name] + marks + [total, avg, grade, result])

    line = f"{reg}  {name:<7} Total={total}  Avg={avg}  Grade={grade:<3} {result}"
    if passed:
        passed_count += 1
        print(line)
    else:
        failed_count += 1
        print(f"{RED}{line}{RESET}")
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", start_color="FFC7CE")

for col, width in zip("ABCDEFGHIJ", [10, 10, 7, 7, 7, 7, 8, 10, 7, 8]):
    ws.column_dimensions[col].width = width

# Summary sheet with chart
summary = wb.create_sheet("Summary")
summary.append(["Result", "Count"])
summary.append(["PASS", passed_count])
summary.append(["FAIL", failed_count])
for cell in summary[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", start_color="4F81BD")

chart = BarChart()
chart.title = "Pass / Fail Summary"
chart.x_axis.title = "Result"
chart.y_axis.title = "Students"
chart.add_data(Reference(summary, min_col=2, min_row=1, max_row=3), titles_from_data=True)
chart.set_categories(Reference(summary, min_col=1, min_row=2, max_row=3))
summary.add_chart(chart, "D2")

wb.save(OUT)

pass_pct = passed_count / len(students) * 100
print(f"\nPass percentage: {pass_pct:.2f} %")
print(f"Report saved to {OUT}  (Passed: {passed_count}, Failed: {failed_count})")
