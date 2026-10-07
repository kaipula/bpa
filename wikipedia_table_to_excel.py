import re
from io import StringIO

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

URL = "https://en.wikipedia.org/wiki/List_of_states_and_union_territories_of_India_by_population"
OUT = "india_population.xlsx"
HEADERS = {"User-Agent": "Mozilla/5.0 (BPA Lab Automation)"}


def flatten(columns):
    names = []
    for col in columns:
        if isinstance(col, tuple):
            parts = [str(p) for p in col if "Unnamed" not in str(p)]
            names.append(" ".join(dict.fromkeys(parts)).strip())
        else:
            names.append(str(col))
    return names


def to_number(series, as_int=False):
    cleaned = series.astype(str).str.replace(r"\[.*?\]", "", regex=True)
    cleaned = cleaned.str.replace(r"[^\d.]", "", regex=True)
    numbers = pd.to_numeric(cleaned, errors="coerce")
    return numbers.astype("Int64") if as_int else numbers


html = requests.get(URL, headers=HEADERS, timeout=20).text
tables = pd.read_html(StringIO(html))
print(f"Tables found on page: {len(tables)}")

target = None
for t in tables:
    t.columns = flatten(t.columns)
    cols = " ".join(t.columns).lower()
    if "population" in cols and ("state" in cols or "union" in cols):
        target = t
        break

state_col = next(c for c in target.columns if "state" in c.lower() or "union" in c.lower())
pop_col = next(c for c in target.columns if "population" in c.lower())
share_col = next(c for c in target.columns if "%" in c or "percent" in c.lower())

df = pd.DataFrame({
    "State": target[state_col].astype(str).str.replace(r"\[.*?\]", "", regex=True).str.strip(),
    "Population": to_number(target[pop_col], as_int=True),
    "Share (%)": to_number(target[share_col]),
})
df = df.dropna(subset=["Population"]).sort_values("Population", ascending=False).head(15)
df.insert(0, "Rank", range(1, len(df) + 1))

df.to_excel(OUT, index=False, sheet_name="Top 15 States")

wb = load_workbook(OUT)
ws = wb.active
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", start_color="1F4E78")
for col, width in zip("ABCD", [7, 22, 16, 11]):
    ws.column_dimensions[col].width = width
for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
    row[0].number_format = "#,##0"
wb.save(OUT)

print(df.head().to_string(index=False))
print(f"\n{len(df)} rows exported to {OUT}")
